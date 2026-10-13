from flask import Flask, render_template, request, jsonify, redirect, url_for, send_file, flash
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, login_user, login_required, logout_user, current_user, UserMixin
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from datetime import datetime, timedelta, date
from zoneinfo import ZoneInfo
from functools import wraps
from collections import defaultdict
from xml.sax.saxutils import escape as xml_escape
from sqlalchemy import func
import pandas as pd
import calendar
import unicodedata
import warnings
import secrets
import random
import time
import io
import os
import re
import json
import requests
import smtplib
from email.mime.text import MIMEText

# ─── FORECAST ────────────────────────────────────────────────────────────────
try:
    from statsmodels.tsa.holtwinters import ExponentialSmoothing
    FORECAST_AVAILABLE = True
except ImportError:
    FORECAST_AVAILABLE = False
    print("⚠️ statsmodels no instalado. Forecast deshabilitado. Ejecuta: pip install statsmodels")

# ─── CONFIGURACIÓN ───────────────────────────────────────────────────────────

# Render (y la mayoría de hostings) definen PORT; en tu PC no existe.
ENV_PRODUCCION = os.environ.get('PORT') is not None

# Zona horaria del negocio: sin esto, en Render "hoy" cambia a las 7 p. m. hora Colombia.
TZ = ZoneInfo(os.environ.get('APP_TZ', 'America/Bogota'))

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)


def _secret_key():
    """SECRET_KEY desde el entorno; en local se genera una y se guarda en instance/."""
    key = os.environ.get('SECRET_KEY')
    if key:
        return key
    ruta = os.path.join(app.instance_path, '.secret_key')
    try:
        os.makedirs(app.instance_path, exist_ok=True)
        if os.path.exists(ruta):
            with open(ruta) as f:
                return f.read().strip()
        key = secrets.token_urlsafe(48)
        with open(ruta, 'w') as f:
            f.write(key)
        return key
    except OSError:
        print('[WARN] SECRET_KEY no configurada: las sesiones se cerrarán en cada reinicio.')
        return secrets.token_urlsafe(48)


def _database_url():
    """DATABASE_URL (Postgres) si existe; si no, SQLite local en instance/inventa.db."""
    url = os.environ.get('DATABASE_URL', '').strip()
    if not url:
        return 'sqlite:///inventa.db'
    if url.startswith('postgres://'):
        url = 'postgresql://' + url[len('postgres://'):]
    if url.startswith('postgresql://'):
        url = 'postgresql+psycopg://' + url[len('postgresql://'):]
    return url


app.config.update(
    SECRET_KEY=_secret_key(),
    SQLALCHEMY_DATABASE_URI=_database_url(),
    SQLALCHEMY_TRACK_MODIFICATIONS=False,
    SQLALCHEMY_ENGINE_OPTIONS={'pool_pre_ping': True},
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=ENV_PRODUCCION,
    REMEMBER_COOKIE_HTTPONLY=True,
    REMEMBER_COOKIE_SAMESITE='Lax',
    REMEMBER_COOKIE_SECURE=ENV_PRODUCCION,
    REMEMBER_COOKIE_DURATION=timedelta(days=30),
    MAX_CONTENT_LENGTH=8 * 1024 * 1024,
)
USA_SQLITE = app.config['SQLALCHEMY_DATABASE_URI'].startswith('sqlite')

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Inicia sesión para continuar'
login_manager.login_message_category = 'info'

# ─── CONFIGURACIÓN DE ALERTAS (solo por variables de entorno) ────────────────
ALERTAS_CONFIG = {
    'telegram_token': os.environ.get('TELEGRAM_BOT_TOKEN', ''),
    'telegram_chat_id': os.environ.get('TELEGRAM_CHAT_ID', ''),
    'smtp_server': os.environ.get('ALERT_SMTP_SERVER', 'smtp.gmail.com'),
    'smtp_port': int(os.environ.get('ALERT_SMTP_PORT', '587')),
    'smtp_user': os.environ.get('ALERT_EMAIL', ''),
    'smtp_password': os.environ.get('ALERT_EMAIL_PASS', ''),
    'alert_dest_email': os.environ.get('ALERT_DEST_EMAIL', ''),
}

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
         'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
ROLES = ('admin', 'vendedor', 'consultor')

# ─── MODELOS ─────────────────────────────────────────────────────────────────

class Usuario(db.Model, UserMixin):
    id              = db.Column(db.Integer, primary_key=True)
    nombre          = db.Column(db.String(100), default='Admin')
    email           = db.Column(db.String(100), unique=True)
    password_hash   = db.Column(db.String(200))
    empresa         = db.Column(db.String(100), default='Mi Empresa')
    capital_inicial = db.Column(db.Float, default=0)
    meta_mensual    = db.Column(db.Float, default=0)
    rol             = db.Column(db.String(50), default='admin')
    configuracion   = db.Column(db.Text, default='{}')
    telefono        = db.Column(db.String(50), default='')
    created_at      = db.Column(db.DateTime, default=datetime.utcnow)

    def set_password(self, pw):
        self.password_hash = generate_password_hash(pw)

    def check_password(self, pw):
        return bool(self.password_hash) and check_password_hash(self.password_hash, pw)

    def get_config(self):
        try:
            return json.loads(self.configuracion or '{}')
        except ValueError:
            return {}

    def set_config(self, config_dict):
        self.configuracion = json.dumps(config_dict)

    @property
    def debe_cambiar_clave(self):
        return bool(self.get_config().get('cambiar_clave'))


class Producto(db.Model):
    id            = db.Column(db.Integer, primary_key=True)
    nombre        = db.Column(db.String(200), nullable=False)
    categoria     = db.Column(db.String(100), default='General')
    precio_venta  = db.Column(db.Float, default=0)
    costo         = db.Column(db.Float, default=0)
    stock_actual  = db.Column(db.Integer, default=0)
    stock_minimo  = db.Column(db.Integer, default=10)
    activo        = db.Column(db.Boolean, default=True)
    creado        = db.Column(db.DateTime, default=datetime.utcnow)
    codigo_barras = db.Column(db.String(100), default='')
    ubicacion     = db.Column(db.String(100), default='')


class VentaDetalle(db.Model):
    id              = db.Column(db.Integer, primary_key=True)
    fecha           = db.Column(db.Date, nullable=False, index=True)
    producto_nombre = db.Column(db.String(200), nullable=False, index=True)
    categoria       = db.Column(db.String(100), default='General')
    cantidad        = db.Column(db.Integer, default=1)
    precio_unitario = db.Column(db.Float, default=0)
    costo_unitario  = db.Column(db.Float, default=0)
    ingreso         = db.Column(db.Float, default=0)
    utilidad        = db.Column(db.Float, default=0)
    subido_en       = db.Column(db.DateTime, default=datetime.utcnow)


class Gasto(db.Model):
    id         = db.Column(db.Integer, primary_key=True)
    concepto   = db.Column(db.String(200))
    monto      = db.Column(db.Float)
    tipo       = db.Column(db.String(50), default='fijo')
    frecuencia = db.Column(db.String(50), default='mensual')
    activo     = db.Column(db.Boolean, default=True)


class AlertaStock(db.Model):
    id              = db.Column(db.Integer, primary_key=True)
    producto_id     = db.Column(db.Integer)
    producto_nombre = db.Column(db.String(200))
    stock_actual    = db.Column(db.Integer)
    stock_minimo    = db.Column(db.Integer)
    tipo            = db.Column(db.String(50))
    enviado_por     = db.Column(db.String(50))
    fecha_envio     = db.Column(db.DateTime, default=datetime.utcnow)
    leido           = db.Column(db.Boolean, default=False)

# ─── UTILIDADES ──────────────────────────────────────────────────────────────

def hoy():
    return datetime.now(TZ).date()


def nombre_mes(d):
    return f'{MESES[d.month - 1].capitalize()} {d.year}'


def error(msg, status=400):
    return jsonify({'success': False, 'error': msg}), status


def parse_num(valor, default=None):
    """Convierte números escritos como 45000, 45.000, 45,000, $45.000 o 12,5 a float."""
    if valor is None:
        return default
    if isinstance(valor, (int, float)):
        return default if pd.isna(valor) else float(valor)
    s = str(valor).strip().replace('$', '').replace(' ', '').replace(' ', '')
    if not s or s.lower() in ('nan', 'none'):
        return default
    if '.' in s and ',' in s:
        # El último separador es el decimal
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif re.fullmatch(r'-?\d{1,3}([.,]\d{3})+', s):
        s = s.replace('.', '').replace(',', '')
    else:
        s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return default


def parse_fecha(valor):
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return None
    if isinstance(valor, pd.Timestamp):
        return valor.date()
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()
    if not s:
        return None
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        pass
    try:
        return pd.to_datetime(s, dayfirst=True).date()
    except (ValueError, TypeError):
        return None


def normalizar_columna(c):
    c = unicodedata.normalize('NFKD', str(c)).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '_', c.strip().lower()).strip('_')


def leer_tabla(file):
    """Lee Excel o CSV y normaliza los nombres de columna."""
    nombre = (file.filename or '').lower()
    if nombre.endswith('.csv'):
        df = pd.read_csv(file, sep=None, engine='python', encoding='utf-8-sig', dtype=str)
    else:
        df = pd.read_excel(file)
    df.columns = [normalizar_columna(c) for c in df.columns]
    return df.dropna(how='all')


def mapear_columnas(df, alias):
    m = {}
    for destino, opciones in alias.items():
        for o in opciones:
            if o in df.columns:
                m[destino] = o
                break
    return m


def valor_celda(row, m, campo):
    if campo not in m:
        return None
    v = row[m[campo]]
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    return v


def texto(v, default=''):
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return default
    s = str(v).strip()
    return s or default


def variacion(actual, anterior):
    """% de cambio; None cuando no hay base de comparación."""
    if not anterior:
        return None
    return round((actual - anterior) / abs(anterior) * 100, 1)


def rango_periodo(periodo=None, inicio=None, fin=None):
    """Devuelve (inicio, fin, etiqueta) para un período con nombre o un rango libre."""
    h = hoy()
    if inicio or fin:
        i = parse_fecha(inicio) or date(2000, 1, 1)
        f = parse_fecha(fin) or h
        if i > f:
            i, f = f, i
        return i, f, f'{i:%d/%m/%Y} – {f:%d/%m/%Y}'
    periodo = periodo or 'mes'
    if periodo == 'dia':
        return h, h, 'Hoy'
    if periodo == 'semana':
        return h - timedelta(days=h.weekday()), h, 'Esta semana'
    if periodo == 'mes':
        return h.replace(day=1), h, nombre_mes(h)
    if periodo == 'mes_anterior':
        f = h.replace(day=1) - timedelta(days=1)
        return f.replace(day=1), f, nombre_mes(f)
    if periodo == 'anio':
        return h.replace(month=1, day=1), h, f'Año {h.year}'
    primera = db.session.query(func.min(VentaDetalle.fecha)).scalar() or h
    return primera, h, 'Todo el historial'


def rango_desde_request():
    return rango_periodo(request.args.get('periodo'), request.args.get('inicio'), request.args.get('fin'))


def totales(inicio, fin):
    r = db.session.query(
        func.coalesce(func.sum(VentaDetalle.ingreso), 0),
        func.coalesce(func.sum(VentaDetalle.utilidad), 0),
        func.coalesce(func.sum(VentaDetalle.cantidad), 0),
        func.count(VentaDetalle.id),
    ).filter(VentaDetalle.fecha.between(inicio, fin)).one()
    ventas, util = float(r[0]), float(r[1])
    return {
        'ventas': round(ventas, 0), 'utilidad': round(util, 0),
        'unidades': int(r[2]), 'transacciones': int(r[3]),
        'margen': round(util / ventas * 100, 1) if ventas else 0,
        'ticket_promedio': round(ventas / r[3], 0) if r[3] else 0,
    }


def por_producto(inicio, fin):
    rows = db.session.query(
        VentaDetalle.producto_nombre,
        func.max(VentaDetalle.categoria),
        func.sum(VentaDetalle.cantidad),
        func.sum(VentaDetalle.ingreso),
        func.sum(VentaDetalle.utilidad),
        func.count(VentaDetalle.id),
    ).filter(VentaDetalle.fecha.between(inicio, fin)) \
     .group_by(VentaDetalle.producto_nombre) \
     .order_by(func.sum(VentaDetalle.ingreso).desc()).all()
    out = []
    for r in rows:
        ing, util = float(r[3] or 0), float(r[4] or 0)
        out.append({
            'nombre': r[0], 'categoria': r[1] or 'General', 'cantidad': int(r[2] or 0),
            'ingreso': round(ing, 0), 'utilidad': round(util, 0), 'transacciones': int(r[5]),
            'margen': round(util / ing * 100, 1) if ing else 0,
        })
    return out


def serie_diaria(inicio, fin):
    rows = db.session.query(
        VentaDetalle.fecha, func.sum(VentaDetalle.ingreso), func.sum(VentaDetalle.utilidad)
    ).filter(VentaDetalle.fecha.between(inicio, fin)).group_by(VentaDetalle.fecha).all()
    mapa = {r[0]: (float(r[1] or 0), float(r[2] or 0)) for r in rows}
    fechas, ingresos, utilidades = [], [], []
    cur = inicio
    while cur <= fin:
        ing, util = mapa.get(cur, (0, 0))
        fechas.append(cur)
        ingresos.append(round(ing, 0))
        utilidades.append(round(util, 0))
        cur += timedelta(days=1)
    return fechas, ingresos, utilidades


def gastos_mensuales():
    return sum(g.monto or 0 for g in Gasto.query.filter_by(activo=True).all())


def estado_stock(p):
    if p.stock_actual <= 0:
        return 'agotado'
    if p.stock_actual <= p.stock_minimo:
        return 'bajo'
    if p.stock_actual <= p.stock_minimo * 2:
        return 'medio'
    return 'ok'


def margen_pct(precio, costo):
    return round((precio - costo) / precio * 100, 1) if precio else 0

# ─── SEGURIDAD ───────────────────────────────────────────────────────────────

def rol_requerido(*roles):
    def deco(f):
        @wraps(f)
        def wrapper(*args, **kwargs):
            if current_user.rol not in roles:
                return error('Tu usuario no tiene permiso para esta acción', 403)
            return f(*args, **kwargs)
        return wrapper
    return deco


PUEDE_EDITAR = ('admin', 'vendedor')


@login_manager.user_loader
def load_user(uid):
    try:
        return db.session.get(Usuario, int(uid))
    except (TypeError, ValueError):
        return None


@login_manager.unauthorized_handler
def no_autorizado():
    if request.path.startswith('/api/'):
        return error('Tu sesión expiró. Vuelve a iniciar sesión.', 401)
    flash('Inicia sesión para continuar', 'info')
    return redirect(url_for('login', next=request.path))


@app.before_request
def proteger_peticiones():
    # Protección CSRF: el navegador no permite que otro sitio envíe este encabezado.
    if request.path.startswith('/api/') and request.method not in ('GET', 'HEAD', 'OPTIONS'):
        if request.headers.get('X-Requested-With') != 'fetch':
            return error('Petición no permitida', 403)
    # Obliga a cambiar la contraseña por defecto antes de usar la app.
    if current_user.is_authenticated and current_user.debe_cambiar_clave:
        libres = ('/cuenta', '/api/cuenta/clave', '/logout', '/static/')
        if not request.path.startswith(libres):
            if request.path.startswith('/api/'):
                return error('Debes cambiar tu contraseña antes de continuar', 403)
            return redirect(url_for('cuenta'))


@app.after_request
def cabeceras_seguridad(resp):
    resp.headers.setdefault('X-Content-Type-Options', 'nosniff')
    resp.headers.setdefault('X-Frame-Options', 'SAMEORIGIN')
    resp.headers.setdefault('Referrer-Policy', 'same-origin')
    return resp


@app.context_processor
def contexto_global():
    return {
        'aviso_disco_temporal': ENV_PRODUCCION and USA_SQLITE,
        'puede_editar': current_user.is_authenticated and current_user.rol in PUEDE_EDITAR,
        'es_admin': current_user.is_authenticated and current_user.rol == 'admin',
    }


# Claves que aparecieron publicadas (README, landing, versiones anteriores): obligan a cambiarlas.
CLAVES_CONOCIDAS = {'inventa2024', 'demo123', 'ver123'}

# Límite simple de intentos de inicio de sesión por IP
_INTENTOS = {}
MAX_INTENTOS, BLOQUEO_SEG = 8, 300


def login_bloqueado(ip):
    fallos, desde = _INTENTOS.get(ip, (0, 0))
    if fallos >= MAX_INTENTOS and time.time() - desde < BLOQUEO_SEG:
        return True
    if time.time() - desde >= BLOQUEO_SEG:
        _INTENTOS.pop(ip, None)
    return False


def registrar_fallo(ip):
    fallos, desde = _INTENTOS.get(ip, (0, time.time()))
    _INTENTOS[ip] = (fallos + 1, desde)

# ─── ALERTAS DE STOCK (Telegram / email, opcionales) ─────────────────────────

def alertas_configuradas():
    c = ALERTAS_CONFIG
    return bool((c['telegram_token'] and c['telegram_chat_id']) or (c['smtp_user'] and c['alert_dest_email']))


def enviar_alerta_telegram(mensaje):
    token, chat_id = ALERTAS_CONFIG['telegram_token'], ALERTAS_CONFIG['telegram_chat_id']
    if not token or not chat_id:
        return False
    try:
        r = requests.post(f"https://api.telegram.org/bot{token}/sendMessage",
                          json={'chat_id': chat_id, 'text': mensaje, 'parse_mode': 'Markdown'}, timeout=5)
        return r.status_code == 200
    except requests.RequestException:
        return False


def enviar_alerta_email(asunto, mensaje):
    c = ALERTAS_CONFIG
    if not c['smtp_user'] or not c['alert_dest_email']:
        return False
    try:
        msg = MIMEText(mensaje, 'html', 'utf-8')
        msg['Subject'] = f"[Inventa] {asunto}"
        msg['From'] = c['smtp_user']
        msg['To'] = c['alert_dest_email']
        with smtplib.SMTP(c['smtp_server'], c['smtp_port'], timeout=10) as server:
            server.starttls()
            server.login(c['smtp_user'], c['smtp_password'])
            server.send_message(msg)
        return True
    except (smtplib.SMTPException, OSError):
        return False


def verificar_y_enviar_alertas_stock():
    if not alertas_configuradas():
        return []
    enviadas = []
    for p in Producto.query.filter_by(activo=True).all():
        if p.stock_actual <= 0:
            tipo, mensaje = 'agotado', f"🚨 *AGOTADO*: {p.nombre} - Stock 0 unidades. ¡Reponer urgentemente!"
        elif p.stock_actual <= p.stock_minimo:
            tipo, mensaje = 'bajo', f"⚠️ *Stock Bajo*: {p.nombre} - {p.stock_actual} uds (mínimo {p.stock_minimo})"
        else:
            continue
        reciente = AlertaStock.query.filter(
            AlertaStock.producto_id == p.id,
            AlertaStock.fecha_envio > datetime.utcnow() - timedelta(hours=24)
        ).first()
        if reciente:
            continue
        ok_tg = enviar_alerta_telegram(mensaje)
        ok_mail = enviar_alerta_email(f"Alerta: {tipo.upper()} - {p.nombre}", mensaje.replace('*', ''))
        for canal, ok in (('telegram', ok_tg), ('email', ok_mail)):
            if ok:
                db.session.add(AlertaStock(producto_id=p.id, producto_nombre=p.nombre, stock_actual=p.stock_actual,
                                           stock_minimo=p.stock_minimo, tipo=tipo, enviado_por=canal))
        enviadas.append(p.nombre)
    if enviadas:
        db.session.commit()
    return enviadas

# ─── FORECAST ────────────────────────────────────────────────────────────────

def generar_forecast(dias=30, seasonal_periods=7):
    if not FORECAST_AVAILABLE:
        return {'error': 'El módulo de pronóstico no está instalado en el servidor'}

    ayer = hoy() - timedelta(days=1)   # el día en curso está incompleto y hundiría la serie
    primera = db.session.query(func.min(VentaDetalle.fecha)).scalar()
    if not primera:
        return {'error': 'Aún no hay ventas registradas'}
    inicio = max(primera, ayer - timedelta(days=119))
    if (ayer - inicio).days + 1 < 14:
        return {'error': f'Se necesitan al menos 14 días de historial (tienes {(ayer - inicio).days + 1})'}

    fechas, serie, _ = serie_diaria(inicio, ayer)
    serie = serie[-90:]
    fechas = fechas[-90:]
    dias_con_ventas = sum(1 for v in serie if v > 0)
    if dias_con_ventas < 10:
        return {'error': f'Solo hay {dias_con_ventas} días con ventas en los últimos meses; se necesitan al menos 10'}
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            # Tendencia amortiguada: con pocas semanas de datos una tendencia lineal promete crecimientos irreales
            model = ExponentialSmoothing(serie, seasonal_periods=seasonal_periods, trend='add', damped_trend=True,
                                         seasonal='add', initialization_method='estimated').fit()
            forecast = [max(0, round(v, 0)) for v in model.forecast(dias).tolist()]
            residuos = list(model.resid)[-30:]
        mae = sum(abs(r) for r in residuos) / len(residuos)
        base = sum(serie[-30:]) / len(serie[-30:]) or 1
        error_pct = mae / base * 100
        return {
            'success': True,
            'forecast': forecast,
            'mae': round(mae, 0),
            'error_pct': round(error_pct, 1),
            'confianza': 'Alta' if error_pct < 20 else 'Media' if error_pct < 40 else 'Baja',
            'dias_analizados': len(serie),
            'historico': {'fechas': [f.strftime('%d/%m') for f in fechas[-30:]], 'valores': serie[-30:]},
        }
    except Exception as e:  # statsmodels puede fallar con series planas o muy cortas
        return {'error': f'No se pudo calcular el pronóstico: {e}'}

# ─── VISTAS ──────────────────────────────────────────────────────────────────

@app.route('/')
def home():
    return redirect(url_for('dashboard' if current_user.is_authenticated else 'login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        ip = request.remote_addr or '?'
        if login_bloqueado(ip):
            flash('Demasiados intentos. Espera 5 minutos e inténtalo de nuevo.', 'error')
            return render_template('login.html'), 429
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        user = Usuario.query.filter(func.lower(Usuario.email) == email).first()
        if user and user.check_password(password):
            _INTENTOS.pop(ip, None)
            if password in CLAVES_CONOCIDAS and not user.debe_cambiar_clave:
                cfg = user.get_config()
                cfg['cambiar_clave'] = True
                user.set_config(cfg)
                db.session.commit()
            login_user(user, remember=bool(request.form.get('recordar')))
            if user.debe_cambiar_clave:
                return redirect(url_for('cuenta'))
            destino = request.args.get('next', '')
            if not destino.startswith('/') or destino.startswith('//'):
                destino = url_for('dashboard')
            return redirect(destino)
        registrar_fallo(ip)
        flash('Correo o contraseña incorrectos', 'error')
    return render_template('login.html')


@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))


@app.route('/dashboard')
@login_required
def dashboard():
    return render_template('dashboard.html')


@app.route('/ventas')
@login_required
def ventas():
    return render_template('ventas.html')


@app.route('/productos')
@login_required
def productos():
    return render_template('productos.html')


@app.route('/proyecciones')
@login_required
def proyecciones():
    return render_template('proyecciones.html')


@app.route('/reportes')
@login_required
def reportes():
    return render_template('reportes.html')


@app.route('/forecast')
@login_required
def forecast_page():
    return render_template('forecast.html')


@app.route('/cuenta')
@login_required
def cuenta():
    return render_template('cuenta.html')

# ─── API: DASHBOARD ──────────────────────────────────────────────────────────

@app.route('/api/dashboard')
@login_required
def api_dashboard():
    h = hoy()
    ini_semana = h - timedelta(days=h.weekday())
    ini_mes = h.replace(day=1)
    # Comparación justa: mismo número de días del mes anterior (1–N vs 1–N)
    fin_mes_ant = ini_mes - timedelta(days=1)
    ini_mes_ant = fin_mes_ant.replace(day=1)
    fin_comparable = min(ini_mes_ant + (h - ini_mes), fin_mes_ant)

    t_hoy = totales(h, h)
    t_ayer = totales(h - timedelta(days=1), h - timedelta(days=1))
    t_semana = totales(ini_semana, h)
    t_mes = totales(ini_mes, h)
    t_mes_ant = totales(ini_mes_ant, fin_comparable)
    t_total = totales(date(2000, 1, 1), h)

    meta = current_user.meta_mensual or 0
    dias_mes = calendar.monthrange(h.year, h.month)[1]
    proyeccion_mes = round(t_mes['ventas'] / h.day * dias_mes, 0) if h.day else 0
    gastos_mes = gastos_mensuales()

    ini30 = h - timedelta(days=29)
    fechas, ingresos, utilidades = serie_diaria(ini30, h)

    top = por_producto(ini_mes, h)[:5]
    alertas = Producto.query.filter(Producto.stock_actual <= Producto.stock_minimo, Producto.activo.is_(True)).count()

    return jsonify({
        'hoy': t_hoy, 'ayer': t_ayer, 'semana': t_semana, 'mes': t_mes,
        'mes_anterior_comparable': t_mes_ant, 'total': t_total,
        'nombre_mes': nombre_mes(h),
        'crecimiento_mensual': variacion(t_mes['ventas'], t_mes_ant['ventas']),
        'meta_mensual': meta,
        'avance_meta': round(t_mes['ventas'] / meta * 100, 1) if meta else None,
        'proyeccion_mes': proyeccion_mes,
        'gastos_mes': round(gastos_mes, 0),
        'ganancia_neta_mes_estimada': round(t_mes['utilidad'] - gastos_mes * h.day / dias_mes, 0),
        'alertas_inventario': alertas,
        'top_productos': top,
        'grafica': {'labels': [f.strftime('%d/%m') for f in fechas], 'ingresos': ingresos, 'utilidades': utilidades},
        'onboarding': {
            'productos': Producto.query.filter_by(activo=True).count(),
            'ventas': VentaDetalle.query.count(),
            'gastos': Gasto.query.filter_by(activo=True).count(),
            'meta': bool(meta),
        },
    })

# ─── API: ANÁLISIS COMPARATIVO MES A MES ─────────────────────────────────────

@app.route('/api/analisis-comparativo')
@login_required
def api_comparativo():
    h = hoy()
    ini_mes = h.replace(day=1)
    fin_mes_ant = ini_mes - timedelta(days=1)
    ini_mes_ant = fin_mes_ant.replace(day=1)
    fin_ant = min(ini_mes_ant + (h - ini_mes), fin_mes_ant)

    actual, anterior = totales(ini_mes, h), totales(ini_mes_ant, fin_ant)
    prods_act = {p['nombre']: p for p in por_producto(ini_mes, h)}
    prods_ant = {p['nombre']: p for p in por_producto(ini_mes_ant, fin_ant)}

    comparables = []
    for nombre, p in prods_act.items():
        if nombre in prods_ant:
            a = prods_ant[nombre]
            comparables.append({
                'nombre': nombre, 'ingreso_actual': p['ingreso'], 'ingreso_anterior': a['ingreso'],
                'variacion': variacion(p['ingreso'], a['ingreso']) or 0,
                'cantidad_actual': p['cantidad'], 'cantidad_anterior': a['cantidad'],
            })
    comparables.sort(key=lambda x: x['variacion'], reverse=True)

    nuevos = sorted([{'nombre': n, 'ingreso': p['ingreso'], 'cantidad': p['cantidad']}
                     for n, p in prods_act.items() if n not in prods_ant], key=lambda x: -x['ingreso'])[:5]
    perdidos = sorted([{'nombre': n, 'ingreso': p['ingreso'], 'cantidad': p['cantidad']}
                       for n, p in prods_ant.items() if n not in prods_act], key=lambda x: -x['ingreso'])[:5]

    var_v = variacion(actual['ventas'], anterior['ventas'])
    return jsonify({
        'mes_actual': {**actual, 'nombre': f'{nombre_mes(h)} (1–{h.day})'},
        'mes_anterior': {**anterior, 'nombre': f'{nombre_mes(ini_mes_ant)} (1–{fin_ant.day})'},
        'variacion': {
            'ventas_pct': var_v,
            'utilidad_pct': variacion(actual['utilidad'], anterior['utilidad']),
            'tendencia': 'positiva' if (var_v or 0) >= 0 else 'negativa',
        },
        'productos_comparables': comparables[:10],
        'productos_nuevos': nuevos,
        'productos_perdidos': perdidos,
    })

# ─── API: FORECAST ───────────────────────────────────────────────────────────

@app.route('/api/forecast')
@login_required
def api_forecast():
    dias = max(7, min(request.args.get('dias', 30, type=int), 90))
    r = generar_forecast(dias=dias)
    if not r.get('success'):
        return jsonify({'success': False, 'error': r.get('error', 'Error generando pronóstico')})
    h = hoy()
    return jsonify({**r, 'fechas': [(h + timedelta(days=i)).strftime('%d/%m') for i in range(dias)]})

# ─── API: PLANTILLAS DE CARGA ────────────────────────────────────────────────

PLANTILLAS = {
    'ventas': (['fecha', 'producto', 'categoria', 'cantidad', 'precio_unitario', 'costo_unitario'], [
        ['2026-01-15', 'Camiseta Básica', 'Ropa', 3, 45000, 20000],
        ['2026-01-15', 'Jeans Premium', 'Ropa', 2, 120000, 60000],
        ['2026-01-16', 'Zapatillas Sport', 'Calzado', 1, 180000, 90000],
    ]),
    'productos': (['nombre', 'categoria', 'precio_venta', 'costo', 'stock_actual', 'stock_minimo'], [
        ['Camiseta Básica', 'Ropa', 45000, 20000, 30, 5],
        ['Jeans Premium', 'Ropa', 120000, 60000, 10, 3],
        ['Zapatillas Sport', 'Calzado', 180000, 90000, 8, 2],
    ]),
}


@app.route('/api/plantilla/<tipo>')
@login_required
def api_plantilla(tipo):
    if tipo not in PLANTILLAS:
        return error('Plantilla no encontrada', 404)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    cols, filas = PLANTILLAS[tipo]
    wb = Workbook()
    ws = wb.active
    ws.title = tipo.capitalize()
    ws.append(cols)
    for f in filas:
        ws.append(f)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[chr(64 + i)].width = 20
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'plantilla_{tipo}.xlsx')

# ─── API: CARGA MASIVA DE PRODUCTOS ──────────────────────────────────────────

@app.route('/api/upload-inventario', methods=['POST'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def upload_inventario():
    file = request.files.get('file')
    if not file:
        return error('No se recibió archivo')
    try:
        df = leer_tabla(file)
    except Exception as e:
        return error(f'No se pudo leer el archivo: {e}')

    m = mapear_columnas(df, {
        'nombre': ['nombre', 'producto', 'articulo', 'item'],
        'precio_venta': ['precio_venta', 'precio', 'pvp', 'valor'],
        'costo': ['costo', 'costo_unitario', 'cost'],
        'categoria': ['categoria', 'category'],
        'stock_actual': ['stock_actual', 'stock', 'existencias', 'cantidad'],
        'stock_minimo': ['stock_minimo', 'minimo'],
        'codigo_barras': ['codigo_barras', 'codigo', 'sku'],
        'ubicacion': ['ubicacion'],
    })
    faltan = [c for c in ('nombre', 'precio_venta') if c not in m]
    if faltan:
        return error(f'Faltan columnas: {", ".join(faltan)}. Encontradas: {", ".join(df.columns)}')

    nuevos = actualizados = 0
    errores = []
    for idx, row in df.iterrows():
        fila = idx + 2
        nombre = texto(valor_celda(row, m, 'nombre'))
        if not nombre:
            errores.append(f'Fila {fila}: sin nombre')
            continue
        precio = parse_num(valor_celda(row, m, 'precio_venta'))
        if precio is None or precio < 0:
            errores.append(f'Fila {fila} ({nombre}): precio inválido')
            continue
        costo = parse_num(valor_celda(row, m, 'costo'), 0)
        datos = {
            'precio_venta': precio, 'costo': max(0, costo),
            'categoria': texto(valor_celda(row, m, 'categoria'), 'General'),
        }
        if 'stock_actual' in m:
            datos['stock_actual'] = int(parse_num(valor_celda(row, m, 'stock_actual'), 0))
        if 'stock_minimo' in m:
            datos['stock_minimo'] = int(parse_num(valor_celda(row, m, 'stock_minimo'), 0))
        for campo in ('codigo_barras', 'ubicacion'):
            if campo in m:
                datos[campo] = texto(valor_celda(row, m, campo))

        p = Producto.query.filter(func.lower(Producto.nombre) == nombre.lower()).first()
        if p:
            for k, v in datos.items():
                setattr(p, k, v)
            p.activo = True
            actualizados += 1
        else:
            datos.setdefault('stock_actual', 0)
            datos.setdefault('stock_minimo', 0)
            db.session.add(Producto(nombre=nombre, **datos))
            nuevos += 1
    db.session.commit()
    verificar_y_enviar_alertas_stock()

    mensaje = f'{nuevos} productos nuevos, {actualizados} actualizados'
    if errores:
        mensaje += f' · {len(errores)} filas con errores'
    return jsonify({'success': True, 'nuevos': nuevos, 'actualizados': actualizados,
                    'errores': errores[:20], 'mensaje': mensaje})

# ─── API: VENTAS ─────────────────────────────────────────────────────────────

def registrar_venta(fecha, nombre, cantidad, precio, costo=None, categoria=None, producto=None):
    """Crea una fila de venta. Si el producto existe, completa costo/categoría y descuenta stock."""
    if producto is None:
        producto = Producto.query.filter(func.lower(Producto.nombre) == nombre.lower(),
                                         Producto.activo.is_(True)).first()
    aviso = None
    if producto:
        nombre = producto.nombre
        if costo is None:
            costo = producto.costo or 0
        categoria = categoria or producto.categoria
        if producto.stock_actual < cantidad and (producto.stock_actual > 0 or producto.stock_minimo > 0):
            aviso = f'{producto.nombre}: vendiste {cantidad} pero el stock registrado era {producto.stock_actual}'
        producto.stock_actual = max(0, (producto.stock_actual or 0) - cantidad)
    costo = costo if costo is not None else 0
    v = VentaDetalle(fecha=fecha, producto_nombre=nombre, categoria=categoria or 'General',
                     cantidad=cantidad, precio_unitario=precio, costo_unitario=costo,
                     ingreso=cantidad * precio, utilidad=cantidad * (precio - costo))
    db.session.add(v)
    return v, aviso, producto is None and costo == 0


@app.route('/api/upload-ventas', methods=['POST'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def upload_ventas():
    file = request.files.get('file')
    if not file:
        return error('No se recibió archivo')
    try:
        df = leer_tabla(file)
    except Exception as e:
        return error(f'No se pudo leer el archivo: {e}')

    m = mapear_columnas(df, {
        'fecha': ['fecha', 'date', 'dia'],
        'producto': ['producto', 'product', 'nombre', 'articulo', 'item'],
        'categoria': ['categoria', 'category'],
        'cantidad': ['cantidad', 'qty', 'quantity', 'unidades'],
        'precio_unitario': ['precio_unitario', 'precio', 'price', 'pvp', 'valor', 'valor_unitario'],
        'costo_unitario': ['costo_unitario', 'costo', 'cost', 'costo_unit'],
    })
    faltan = [c for c in ('producto', 'cantidad') if c not in m]
    if faltan:
        return error(f'Faltan columnas: {", ".join(faltan)}. Encontradas: {", ".join(df.columns)}')

    h = hoy()
    insertadas, sin_costo, errores, avisos = 0, 0, [], []
    for idx, row in df.iterrows():
        fila = idx + 2
        prod = texto(valor_celda(row, m, 'producto'))
        if not prod:
            errores.append(f'Fila {fila}: sin producto')
            continue
        cant = parse_num(valor_celda(row, m, 'cantidad'))
        if not cant or cant <= 0:
            errores.append(f'Fila {fila} ({prod}): cantidad inválida')
            continue
        catalogo = Producto.query.filter(func.lower(Producto.nombre) == prod.lower(),
                                         Producto.activo.is_(True)).first()
        precio = parse_num(valor_celda(row, m, 'precio_unitario'))
        if precio is None:
            if not catalogo:
                errores.append(f'Fila {fila} ({prod}): sin precio y no está en el catálogo')
                continue
            precio = catalogo.precio_venta
        fecha = parse_fecha(valor_celda(row, m, 'fecha')) if 'fecha' in m else None
        if 'fecha' in m and valor_celda(row, m, 'fecha') is not None and fecha is None:
            errores.append(f'Fila {fila} ({prod}): fecha no reconocida')
            continue
        fecha = fecha or h
        if fecha > h:
            errores.append(f'Fila {fila} ({prod}): fecha futura {fecha:%d/%m/%Y}')
            continue
        costo = parse_num(valor_celda(row, m, 'costo_unitario'))
        _, aviso, falta_costo = registrar_venta(fecha, prod, int(cant), precio, costo,
                                                texto(valor_celda(row, m, 'categoria')) or None, catalogo)
        if aviso:
            avisos.append(aviso)
        if falta_costo:
            sin_costo += 1
        insertadas += 1

    db.session.commit()
    verificar_y_enviar_alertas_stock()
    return jsonify({'success': True, 'insertadas': insertadas, 'errores': errores[:20],
                    'total_errores': len(errores), 'sin_costo': sin_costo, 'avisos': avisos[:10]})


@app.route('/api/ventas', methods=['POST'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def api_venta_crear():
    d = request.get_json(silent=True) or {}
    items = d.get('items') or []
    if not items:
        return error('Agrega al menos un producto a la venta')
    fecha = parse_fecha(d.get('fecha')) or hoy()
    if fecha > hoy():
        return error('La fecha de la venta no puede ser futura')

    avisos, creadas = [], []
    for i, it in enumerate(items, 1):
        producto = None
        if it.get('producto_id'):
            producto = db.session.get(Producto, int(it['producto_id']))
            if not producto or not producto.activo:
                db.session.rollback()
                return error(f'Producto {i}: no existe en el catálogo')
        nombre = producto.nombre if producto else texto(it.get('nombre'))
        if not nombre:
            db.session.rollback()
            return error(f'Producto {i}: escribe el nombre o elígelo del catálogo')
        cantidad = parse_num(it.get('cantidad'))
        if not cantidad or cantidad <= 0 or cantidad != int(cantidad):
            db.session.rollback()
            return error(f'{nombre}: la cantidad debe ser un número entero mayor que 0')
        precio = parse_num(it.get('precio_unitario'))
        if precio is None:
            precio = producto.precio_venta if producto else None
        if precio is None or precio < 0:
            db.session.rollback()
            return error(f'{nombre}: indica el precio de venta')
        costo = parse_num(it.get('costo_unitario'))
        if costo is not None and costo < 0:
            db.session.rollback()
            return error(f'{nombre}: el costo no puede ser negativo')
        v, aviso, _ = registrar_venta(fecha, nombre, int(cantidad), precio, costo,
                                      texto(it.get('categoria')) or None, producto)
        creadas.append(v)
        if aviso:
            avisos.append(aviso)

    db.session.commit()
    verificar_y_enviar_alertas_stock()
    total = sum(v.ingreso for v in creadas)
    ganancia = sum(v.utilidad for v in creadas)
    return jsonify({'success': True, 'ids': [v.id for v in creadas], 'total': round(total, 0),
                    'ganancia': round(ganancia, 0), 'avisos': avisos})


@app.route('/api/ventas/lista')
@login_required
def api_ventas_lista():
    inicio, fin, etiqueta = rango_desde_request()
    limite = max(1, min(request.args.get('limite', 300, type=int), 2000))
    q = VentaDetalle.query.filter(VentaDetalle.fecha.between(inicio, fin))
    busqueda = request.args.get('q', '').strip()
    if busqueda:
        q = q.filter(VentaDetalle.producto_nombre.ilike(f'%{busqueda}%'))
    filas = q.order_by(VentaDetalle.fecha.desc(), VentaDetalle.id.desc()).limit(limite).all()
    return jsonify({'periodo': etiqueta, 'ventas': [{
        'id': v.id, 'fecha': v.fecha.isoformat(), 'producto': v.producto_nombre, 'categoria': v.categoria,
        'cantidad': v.cantidad, 'precio_unitario': v.precio_unitario, 'costo_unitario': v.costo_unitario,
        'ingreso': round(v.ingreso, 0), 'utilidad': round(v.utilidad, 0),
    } for v in filas]})


@app.route('/api/ventas/<int:vid>', methods=['DELETE'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def api_venta_eliminar(vid):
    v = db.session.get(VentaDetalle, vid)
    if not v:
        return error('Venta no encontrada', 404)
    p = Producto.query.filter_by(nombre=v.producto_nombre, activo=True).first()
    if p:
        p.stock_actual = (p.stock_actual or 0) + (v.cantidad or 0)
    db.session.delete(v)
    db.session.commit()
    return jsonify({'success': True})


@app.route('/api/ventas/analisis')
@login_required
def api_ventas_analisis():
    inicio, fin, etiqueta = rango_desde_request()
    lista = por_producto(inicio, fin)
    cat = defaultdict(float)
    for p in lista:
        cat[p['categoria']] += p['ingreso']
    vendidos = {p['nombre'] for p in lista}
    sin_ventas = [p.nombre for p in Producto.query.filter_by(activo=True).all() if p.nombre not in vendidos]

    return jsonify({
        'periodo': etiqueta,
        'rango': {'inicio': inicio.isoformat(), 'fin': fin.isoformat()},
        'totales': totales(inicio, fin),
        'productos': lista,
        # Estrella = el que más ganancia deja; no siempre es el que más factura.
        'estrella': max(lista, key=lambda p: p['utilidad']) if lista else None,
        'mas_vendido': max(lista, key=lambda p: p['cantidad']) if lista else None,
        'peor': min(lista, key=lambda p: p['utilidad']) if len(lista) > 1 else None,
        'sin_ventas': sin_ventas,
        'top5': lista[:5],
        'categorias': [{'nombre': k, 'ingreso': round(v, 0)} for k, v in sorted(cat.items(), key=lambda x: -x[1])],
    })

# ─── API: PRODUCTOS ──────────────────────────────────────────────────────────

def datos_producto(d, parcial=False):
    """Valida el JSON de un producto. Devuelve (datos, mensaje_error)."""
    out = {}
    if not parcial or 'nombre' in d:
        nombre = texto(d.get('nombre'))
        if not nombre:
            return None, 'El nombre es obligatorio'
        out['nombre'] = nombre[:200]
    if not parcial or 'categoria' in d:
        out['categoria'] = texto(d.get('categoria'), 'General')[:100]
    for campo, etiqueta in (('precio_venta', 'precio de venta'), ('costo', 'costo')):
        if not parcial or campo in d:
            v = parse_num(d.get(campo), 0 if campo == 'costo' else None)
            if v is None or v < 0:
                return None, f'El {etiqueta} no es válido'
            out[campo] = v
    for campo in ('stock_actual', 'stock_minimo'):
        if not parcial or campo in d:
            v = parse_num(d.get(campo), 0)
            if v is None or v < 0:
                return None, 'El stock no puede ser negativo'
            out[campo] = int(v)
    return out, None


@app.route('/api/productos', methods=['GET'])
@login_required
def api_productos_get():
    h = hoy()
    ventas30 = {p['nombre']: p for p in por_producto(h - timedelta(days=29), h)}
    ps = Producto.query.filter_by(activo=True).order_by(Producto.nombre).all()
    return jsonify([{
        'id': p.id, 'nombre': p.nombre, 'categoria': p.categoria,
        'precio_venta': p.precio_venta, 'costo': p.costo,
        'stock_actual': p.stock_actual, 'stock_minimo': p.stock_minimo,
        'margen': margen_pct(p.precio_venta, p.costo),
        'ganancia_unidad': round((p.precio_venta or 0) - (p.costo or 0), 0),
        'estado': estado_stock(p),
        'vendidas_30d': ventas30.get(p.nombre, {}).get('cantidad', 0),
        'utilidad_30d': ventas30.get(p.nombre, {}).get('utilidad', 0),
    } for p in ps])


@app.route('/api/productos', methods=['POST'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def api_productos_add():
    datos, msg = datos_producto(request.get_json(silent=True) or {})
    if msg:
        return error(msg)
    existente = Producto.query.filter(func.lower(Producto.nombre) == datos['nombre'].lower()).first()
    if existente and existente.activo:
        return error(f'Ya existe un producto llamado "{existente.nombre}"')
    if existente:  # estaba eliminado: se reactiva con los datos nuevos
        for k, v in datos.items():
            setattr(existente, k, v)
        existente.activo = True
        p = existente
    else:
        p = Producto(**datos)
        db.session.add(p)
    db.session.commit()
    return jsonify({'success': True, 'id': p.id})


@app.route('/api/productos/<int:pid>', methods=['PUT'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def api_productos_update(pid):
    p = db.session.get(Producto, pid)
    if not p or not p.activo:
        return error('Producto no encontrado', 404)
    datos, msg = datos_producto(request.get_json(silent=True) or {}, parcial=True)
    if msg:
        return error(msg)
    nuevo_nombre = datos.get('nombre')
    if nuevo_nombre and nuevo_nombre != p.nombre:
        choque = Producto.query.filter(func.lower(Producto.nombre) == nuevo_nombre.lower(),
                                       Producto.id != p.id, Producto.activo.is_(True)).first()
        if choque:
            return error(f'Ya existe un producto llamado "{choque.nombre}"')
        # Las ventas guardan el nombre: se renombran para no perder el historial del producto
        VentaDetalle.query.filter_by(producto_nombre=p.nombre).update({'producto_nombre': nuevo_nombre})
    for k, v in datos.items():
        setattr(p, k, v)
    db.session.commit()
    return jsonify({'success': True})


@app.route('/api/productos/<int:pid>', methods=['DELETE'])
@login_required
@rol_requerido(*PUEDE_EDITAR)
def api_productos_delete(pid):
    p = db.session.get(Producto, pid)
    if not p:
        return error('Producto no encontrado', 404)
    p.activo = False
    db.session.commit()
    return jsonify({'success': True})

# ─── API: INVENTARIO ─────────────────────────────────────────────────────────

@app.route('/api/inventario')
@login_required
def api_inventario():
    alertas, sugerencias, promover, resultado = [], [], [], []
    for p in Producto.query.filter_by(activo=True).order_by(Producto.nombre).all():
        margen = margen_pct(p.precio_venta, p.costo)
        estado = estado_stock(p)
        # stock_minimo = 0 significa "no controlo stock de este producto"
        controla = (p.stock_minimo or 0) > 0
        if controla and estado in ('agotado', 'bajo'):
            alertas.append({'producto': p.nombre, 'stock_actual': p.stock_actual,
                            'stock_minimo': p.stock_minimo, 'estado': estado})
            cant = max(p.stock_minimo * 2 - p.stock_actual, 1)
            sugerencias.append({'producto': p.nombre, 'cantidad_sugerida': cant, 'costo_unitario': p.costo,
                                'inversion_estimada': round(cant * (p.costo or 0), 0)})
        if controla and estado == 'ok' and p.stock_actual > p.stock_minimo * 3 and margen > 20:
            promover.append({'producto': p.nombre, 'stock_actual': p.stock_actual, 'margen': margen,
                             'precio_venta': p.precio_venta})
        resultado.append({'id': p.id, 'nombre': p.nombre, 'categoria': p.categoria, 'precio_venta': p.precio_venta,
                          'costo': p.costo, 'stock_actual': p.stock_actual, 'stock_minimo': p.stock_minimo,
                          'estado': estado if controla else 'sin_control', 'margen': margen})
    return jsonify({'productos': resultado, 'alertas': alertas, 'sugerencias': sugerencias,
                    'promover': promover, 'total_alertas': len(alertas)})

# ─── API: GASTOS ─────────────────────────────────────────────────────────────

@app.route('/api/gastos', methods=['GET'])
@login_required
def api_gastos_get():
    gs = Gasto.query.filter_by(activo=True).all()
    return jsonify({
        'gastos': [{'id': g.id, 'concepto': g.concepto, 'monto': g.monto, 'tipo': g.tipo,
                    'frecuencia': g.frecuencia} for g in gs],
        'total_fijo': sum(g.monto for g in gs if g.tipo == 'fijo'),
        'total_variable': sum(g.monto for g in gs if g.tipo == 'variable'),
    })


@app.route('/api/gastos', methods=['POST'])
@login_required
@rol_requerido('admin')
def api_gastos_add():
    d = request.get_json(silent=True) or {}
    concepto = texto(d.get('concepto'))
    monto = parse_num(d.get('monto'))
    if not concepto or not monto or monto <= 0:
        return error('Escribe el concepto y un monto mayor que 0')
    tipo = d.get('tipo') if d.get('tipo') in ('fijo', 'variable') else 'fijo'
    g = Gasto(concepto=concepto[:200], monto=monto, tipo=tipo, frecuencia='mensual')
    db.session.add(g)
    db.session.commit()
    return jsonify({'success': True, 'id': g.id})


@app.route('/api/gastos/<int:gid>', methods=['DELETE'])
@login_required
@rol_requerido('admin')
def api_gastos_delete(gid):
    g = db.session.get(Gasto, gid)
    if not g:
        return error('Gasto no encontrado', 404)
    g.activo = False
    db.session.commit()
    return jsonify({'success': True})

# ─── API: PROYECCIONES ───────────────────────────────────────────────────────

@app.route('/api/proyeccion/calcular', methods=['POST'])
@login_required
def api_proyeccion():
    d = request.get_json(silent=True) or {}
    capital = max(0, parse_num(d.get('capital'), 0))
    gastos_mes = parse_num(d.get('gastos_fijos_mes'))
    if gastos_mes is None:
        gastos_mes = gastos_mensuales()
    ventas_dia = parse_num(d.get('ventas_diarias_estimadas'), 10)
    if ventas_dia is None or ventas_dia < 0:
        return error('Las ventas diarias estimadas no son válidas')

    # Mezcla de productos: si hay ventas recientes se pondera por lo que realmente se vende;
    # si no, cada producto del catálogo pesa igual.
    h = hoy()
    reales = por_producto(h - timedelta(days=59), h)
    if reales and sum(p['cantidad'] for p in reales):
        tqty = sum(p['cantidad'] for p in reales)
        tp = sum(p['ingreso'] for p in reales)
        tc = tp - sum(p['utilidad'] for p in reales)
        base = 'ventas reales de los últimos 60 días'
    else:
        ps = Producto.query.filter_by(activo=True).all()
        if not ps:
            return error('Agrega productos al catálogo o registra ventas primero')
        tqty, tp, tc = len(ps), sum(p.precio_venta for p in ps), sum(p.costo for p in ps)
        base = 'promedio simple del catálogo'

    if tp <= 0:
        return error('Los precios de venta están en 0')
    pprom, cprom = tp / tqty, tc / tqty
    margen_u = pprom - cprom
    if margen_u <= 0:
        return error('En promedio el costo es igual o mayor al precio: cada venta da pérdida')

    gdia = gastos_mes / 30
    uq_eq_dia = gdia / margen_u
    ing_dia = ventas_dia * pprom
    util_dia = ventas_dia * margen_u - gdia
    dias_rec = 0 if capital == 0 else (round(capital / util_dia, 0) if util_dia > 0 else -1)

    proyecciones, acum, mes_utilidad = [], -capital, (0 if capital == 0 and util_dia > 0 else None)
    for mes in range(1, 13):
        util_mes = util_dia * 30
        acum += util_mes
        proyecciones.append({'mes': mes, 'ingreso': round(ing_dia * 30, 0),
                             'costos': round((cprom * ventas_dia + gdia) * 30, 0),
                             'utilidad': round(util_mes, 0), 'capital_acumulado': round(acum, 0)})
        if mes_utilidad is None and acum >= 0:
            mes_utilidad = mes

    return jsonify({
        'base_calculo': base,
        'precio_promedio': round(pprom, 0), 'costo_promedio': round(cprom, 0),
        'margen_contribucion': round(margen_u, 0), 'margen_porcentaje': round(margen_u / pprom * 100, 1),
        'gastos_fijos_dia': round(gdia, 0), 'unidades_equilibrio_dia': round(uq_eq_dia, 1),
        'ventas_equilibrio_dia': round(uq_eq_dia * pprom, 0), 'ingreso_diario_estimado': round(ing_dia, 0),
        'utilidad_diaria_estimada': round(util_dia, 0), 'dias_recuperacion_capital': dias_rec,
        'mes_inicio_utilidades': mes_utilidad, 'proyecciones_mensuales': proyecciones,
    })

# ─── API: REPORTES ───────────────────────────────────────────────────────────

def resumen_reporte(inicio, fin):
    t = totales(inicio, fin)
    dias = (fin - inicio).days + 1
    gastos_periodo = gastos_mensuales() * dias / 30
    return t, round(gastos_periodo, 0), round(t['utilidad'] - gastos_periodo, 0)


@app.route('/api/reporte-data')
@login_required
def api_reporte_data():
    inicio, fin, etiqueta = rango_desde_request()
    t, gastos_periodo, neta = resumen_reporte(inicio, fin)
    return jsonify({
        'periodo': etiqueta, 'inicio': inicio.isoformat(), 'fin': fin.isoformat(),
        'ventas_total': t['ventas'], 'utilidad_total': t['utilidad'], 'margen': t['margen'],
        'total_transacciones': t['transacciones'], 'unidades': t['unidades'],
        'ticket_promedio': t['ticket_promedio'],
        'gastos_periodo': gastos_periodo, 'ganancia_neta_estimada': neta,
        'top_productos': por_producto(inicio, fin)[:10],
    })


@app.route('/api/reporte/excel')
@login_required
def api_reporte_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    inicio, fin, etiqueta = rango_desde_request()
    t, gastos_periodo, neta = resumen_reporte(inicio, fin)
    wb = Workbook()
    hdr_font = Font(bold=True, color='FFFFFF', size=11)
    fill_a = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
    fill_b = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    MONEDA = '"$"#,##0'

    def hoja(ws, headers, fill, filas, anchos, monedas=()):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hdr_font, fill, Alignment(horizontal='center')
        for f in filas:
            ws.append(f)
        for i, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in monedas:
            for celda in ws[get_column_letter(col)][1:]:
                celda.number_format = MONEDA
        ws.freeze_panes = 'A2'

    ws0 = wb.active
    ws0.title = 'Resumen'
    ws0.append([f'{current_user.empresa} — Reporte de ventas'])
    ws0['A1'].font = Font(bold=True, size=14)
    ws0.append([f'Período: {etiqueta} ({inicio:%d/%m/%Y} – {fin:%d/%m/%Y})'])
    ws0.append([])
    for nombre, valor, moneda in (
        ('Ventas', t['ventas'], True), ('Ganancia bruta', t['utilidad'], True),
        ('Margen %', t['margen'], False), ('Transacciones', t['transacciones'], False),
        ('Unidades vendidas', t['unidades'], False), ('Ticket promedio', t['ticket_promedio'], True),
        ('Gastos del período (estimado)', gastos_periodo, True),
        ('Ganancia después de gastos (estimada)', neta, True),
    ):
        ws0.append([nombre, valor])
        if moneda:
            ws0.cell(row=ws0.max_row, column=2).number_format = MONEDA
    ws0.column_dimensions['A'].width = 40
    ws0.column_dimensions['B'].width = 20

    ventas_q = VentaDetalle.query.filter(VentaDetalle.fecha.between(inicio, fin)) \
        .order_by(VentaDetalle.fecha.desc(), VentaDetalle.id.desc()).all()
    hoja(wb.create_sheet('Ventas'),
         ['Fecha', 'Producto', 'Categoría', 'Cantidad', 'Precio', 'Costo', 'Ingreso', 'Ganancia'], fill_a,
         [[v.fecha.strftime('%Y-%m-%d'), v.producto_nombre, v.categoria, v.cantidad, v.precio_unitario,
           v.costo_unitario, v.ingreso, v.utilidad] for v in ventas_q],
         [12, 32, 16, 10, 14, 14, 16, 16], monedas=(5, 6, 7, 8))

    hoja(wb.create_sheet('Por producto'),
         ['Producto', 'Categoría', 'Unidades', 'Ingresos', 'Ganancia', 'Margen %'], fill_b,
         [[p['nombre'], p['categoria'], p['cantidad'], p['ingreso'], p['utilidad'], p['margen']]
          for p in por_producto(inicio, fin)],
         [32, 16, 12, 16, 16, 12], monedas=(4, 5))

    prods = Producto.query.filter_by(activo=True).order_by(Producto.nombre).all()
    hoja(wb.create_sheet('Productos'),
         ['Producto', 'Categoría', 'Precio', 'Costo', 'Margen %', 'Stock', 'Stock mín.'], fill_a,
         [[p.nombre, p.categoria, p.precio_venta, p.costo, margen_pct(p.precio_venta, p.costo),
           p.stock_actual, p.stock_minimo] for p in prods],
         [32, 16, 14, 14, 12, 10, 12], monedas=(3, 4))

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'reporte_{inicio:%Y%m%d}_{fin:%Y%m%d}.xlsx')


@app.route('/api/reporte/pdf')
@login_required
def api_reporte_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.units import cm

    inicio, fin, etiqueta = rango_desde_request()
    t, gastos_periodo, neta = resumen_reporte(inicio, fin)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title=f'Reporte {etiqueta}', author=current_user.empresa or 'Inventa')
    styles = getSampleStyleSheet()
    AZUL, OSCURO, CLARO = colors.HexColor('#1D4ED8'), colors.HexColor('#0F172A'), colors.HexColor('#F1F5F9')
    t_style = ParagraphStyle('t', parent=styles['Heading1'], fontSize=20, textColor=OSCURO, spaceAfter=2)
    sub = ParagraphStyle('s', parent=styles['Normal'], textColor=colors.HexColor('#475569'))
    h2 = ParagraphStyle('h2', parent=styles['Heading2'], fontSize=13, textColor=AZUL, spaceBefore=10, spaceAfter=6)
    nota = ParagraphStyle('n', parent=styles['Normal'], fontSize=8, textColor=colors.HexColor('#64748b'))

    def tabla(data, anchos, color):
        tb = Table(data, colWidths=anchos, repeatRows=1)
        tb.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), color), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [CLARO, colors.white]),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CBD5E1')), ('PADDING', (0, 0), (-1, -1), 6),
            ('ALIGN', (1, 1), (-1, -1), 'RIGHT'),
        ]))
        return tb

    m = lambda v: f'${v:,.0f}'.replace(',', '.')
    story = [
        Paragraph(xml_escape(current_user.empresa or 'Mi negocio'), t_style),
        Paragraph(f'Reporte de ventas · {xml_escape(etiqueta)} ({inicio:%d/%m/%Y} – {fin:%d/%m/%Y})', sub),
        Paragraph(f'Generado el {datetime.now(TZ):%d/%m/%Y %H:%M}', sub),
        Paragraph('Resumen', h2),
        tabla([['Indicador', 'Valor'],
               ['Ventas', m(t['ventas'])], ['Ganancia bruta', m(t['utilidad'])],
               ['Margen', f"{t['margen']}%"], ['Transacciones', str(t['transacciones'])],
               ['Unidades vendidas', str(t['unidades'])], ['Ticket promedio', m(t['ticket_promedio'])],
               ['Gastos del período (estimado)', m(gastos_periodo)],
               ['Ganancia después de gastos (estimada)', m(neta)]], [10 * cm, 7 * cm], AZUL),
        Spacer(1, 0.15 * cm),
        Paragraph('Los gastos del período se estiman prorrateando los gastos mensuales registrados.', nota),
        Paragraph('Productos del período', h2),
    ]
    tops = por_producto(inicio, fin)[:15]
    if tops:
        story.append(tabla([['Producto', 'Unid.', 'Ingresos', 'Ganancia', 'Margen']] +
                           [[p['nombre'][:38], str(p['cantidad']), m(p['ingreso']), m(p['utilidad']),
                             f"{p['margen']}%"] for p in tops],
                           [6.6 * cm, 1.6 * cm, 3.2 * cm, 3.2 * cm, 2.4 * cm], OSCURO))
    else:
        story.append(Paragraph('No hay ventas en este período.', styles['Normal']))

    alertas = Producto.query.filter(Producto.stock_actual <= Producto.stock_minimo, Producto.stock_minimo > 0,
                                    Producto.activo.is_(True)).all()
    if alertas:
        story += [Paragraph('Productos por reponer', h2),
                  tabla([['Producto', 'Stock', 'Mínimo', 'Estado']] +
                        [[p.nombre[:40], str(p.stock_actual), str(p.stock_minimo),
                          'AGOTADO' if p.stock_actual <= 0 else 'BAJO'] for p in alertas],
                        [8 * cm, 3 * cm, 3 * cm, 3 * cm], colors.HexColor('#DC2626'))]

    doc.build(story)
    buf.seek(0)
    return send_file(buf, mimetype='application/pdf', as_attachment=True,
                     download_name=f'reporte_{inicio:%Y%m%d}_{fin:%Y%m%d}.pdf')

# ─── API: ALERTAS ────────────────────────────────────────────────────────────

@app.route('/api/alertas')
@login_required
def api_alertas_lista():
    alertas = AlertaStock.query.order_by(AlertaStock.fecha_envio.desc()).limit(50).all()
    return jsonify([{'id': a.id, 'producto': a.producto_nombre, 'stock_actual': a.stock_actual,
                     'stock_minimo': a.stock_minimo, 'tipo': a.tipo,
                     'fecha': a.fecha_envio.strftime('%Y-%m-%d %H:%M'), 'leido': a.leido} for a in alertas])


@app.route('/api/alertas/marcar-leida/<int:alerta_id>', methods=['POST'])
@login_required
def api_alerta_marcar_leida(alerta_id):
    alerta = db.session.get(AlertaStock, alerta_id)
    if not alerta:
        return error('Alerta no encontrada', 404)
    alerta.leido = True
    db.session.commit()
    return jsonify({'success': True})

# ─── API: ANÁLISIS INTELIGENTE ───────────────────────────────────────────────

@app.route('/api/analisis-inteligente')
@login_required
def api_analisis_inteligente():
    h = hoy()
    ini60 = h - timedelta(days=59)
    rows = db.session.query(
        VentaDetalle.producto_nombre,
        func.sum(VentaDetalle.cantidad), func.sum(VentaDetalle.ingreso),
        func.sum(VentaDetalle.utilidad), func.min(VentaDetalle.fecha),
    ).filter(VentaDetalle.fecha >= ini60).group_by(VentaDetalle.producto_nombre).all()

    ventas = {}
    for r in rows:
        dias = max(7, (h - r[4]).days + 1)   # velocidad desde la primera venta, no siempre 60 días
        ventas[r[0]] = {'qty': int(r[1]), 'ing': float(r[2]), 'util': float(r[3]), 'velocidad': round(int(r[1]) / dias, 2)}

    productos = Producto.query.filter_by(activo=True).all()
    sacar, promover, mantener, sin_ventas = [], [], [], []
    for p in productos:
        mg = margen_pct(p.precio_venta, p.costo)
        v = ventas.get(p.nombre)
        base = {'nombre': p.nombre, 'categoria': p.categoria, 'stock_actual': p.stock_actual,
                'precio_venta': p.precio_venta, 'costo': p.costo, 'margen': mg}
        if not v:
            sin_ventas.append({**base, 'razon': 'Sin ventas en los últimos 60 días'})
            continue
        vel = v['velocidad']
        dias_stock = p.stock_actual / vel if vel > 0 else 9999
        desc = min(30, max(10, int((dias_stock - 30) / 10) * 5)) if dias_stock > 30 else 10
        precio_promo = round(max((p.costo or 0) * 1.05, p.precio_venta * (1 - desc / 100)), -2)
        item = {**base, 'velocidad': vel, 'dias_stock_restante': round(min(dias_stock, 9999), 0),
                'precio_promo': precio_promo, 'descuento_pct': desc}
        if dias_stock > 90 and mg < 15:
            item['razon'] = f'Mucho stock ({int(dias_stock)} días) y margen bajo ({mg}%)'
            sacar.append(item)
        elif dias_stock > 45 and vel < 0.5:
            item['razon'] = f'Se vende despacio ({vel} por día)'
            promover.append(item)
        elif mg > 30 and vel < 1.0 and p.stock_actual > p.stock_minimo * 2:
            item['razon'] = 'Buen margen pero se vende poco'
            promover.append(item)
        else:
            mantener.append(item)

    return jsonify({
        'sacar': sacar, 'promover': promover, 'mantener': mantener, 'sin_ventas': sin_ventas,
        'resumen': {
            'total_productos': len(productos),
            'productos_con_ventas': len([p for p in productos if p.nombre in ventas]),
            'ventas_60_dias': round(sum(v['ing'] for v in ventas.values()), 0),
            'utilidad_60_dias': round(sum(v['util'] for v in ventas.values()), 0),
            'criticos': len(sacar) + len(sin_ventas),
            'oportunidades': len(promover),
        },
    })

# ─── API: DATOS DE DEMOSTRACIÓN Y LIMPIEZA (solo administrador) ──────────────

@app.route('/api/demo-data', methods=['POST'])
@login_required
@rol_requerido('admin')
def api_demo_data():
    d = request.get_json(silent=True) or {}
    if VentaDetalle.query.count() and d.get('confirmar') != 'BORRAR':
        return error('Ya tienes ventas registradas. Cargar la demo las borra: escribe BORRAR para confirmar.', 409)

    rnd = random.Random(42)
    VentaDetalle.query.delete()
    Producto.query.delete()
    Gasto.query.delete()
    db.session.commit()

    demo_prods = [
        ('Camiseta Basica', 'Ropa', 45000, 20000, 150, 20),
        ('Jeans Premium', 'Ropa', 120000, 60000, 80, 15),
        ('Zapatillas Sport', 'Calzado', 180000, 90000, 60, 10),
        ('Mochila Urbana', 'Accesorios', 95000, 45000, 45, 8),
        ('Audifonos BT', 'Electronica', 85000, 40000, 30, 5),
        ('Reloj Casual', 'Accesorios', 65000, 30000, 25, 5),
        ('Sudadera Hoodie', 'Ropa', 89000, 42000, 90, 15),
        ('Gorra Snapback', 'Accesorios', 35000, 15000, 120, 10),
        ('Lentes de Sol', 'Accesorios', 55000, 25000, 40, 8),
        ('Botas Tacticas', 'Calzado', 220000, 120000, 20, 5),
    ]
    prods = [Producto(nombre=n, categoria=c, precio_venta=pv, costo=co, stock_actual=sa, stock_minimo=sm)
             for n, c, pv, co, sa, sm in demo_prods]
    db.session.add_all(prods)
    db.session.add_all([
        Gasto(concepto='Arriendo', monto=1500000, tipo='fijo'),
        Gasto(concepto='Servicios', monto=250000, tipo='fijo'),
        Gasto(concepto='Nomina', monto=2500000, tipo='fijo'),
        Gasto(concepto='Marketing', monto=300000, tipo='variable'),
        Gasto(concepto='Transporte', monto=150000, tipo='variable'),
    ])
    h = hoy()
    vendidos = defaultdict(int)
    ventas_demo = []
    for dd in range(60, -1, -1):
        fecha = h - timedelta(days=dd)
        for _ in range(rnd.randint(5, 15)):
            n, c, pv, co, _, _ = rnd.choice(demo_prods)
            cant = rnd.randint(1, 8)
            precio = round(pv * rnd.uniform(0.92, 1.08), -2)
            vendidos[n] += cant
            ventas_demo.append(VentaDetalle(fecha=fecha, producto_nombre=n, categoria=c, cantidad=cant,
                                            precio_unitario=precio, costo_unitario=co,
                                            ingreso=cant * precio, utilidad=cant * (precio - co)))
    db.session.add_all(ventas_demo)
    for p in prods:
        p.stock_actual = max(0, p.stock_actual - vendidos[p.nombre])
    db.session.commit()
    return jsonify({'success': True, 'mensaje': f'Demo cargada: {len(prods)} productos y {len(ventas_demo)} ventas'})


@app.route('/api/datos', methods=['DELETE'])
@login_required
@rol_requerido('admin')
def api_limpiar_datos():
    tipo = request.args.get('tipo', 'ventas')
    if request.args.get('confirmar') != 'BORRAR':
        return error('Escribe BORRAR para confirmar')
    if tipo == 'ventas':
        VentaDetalle.query.delete()
    elif tipo == 'todo':
        VentaDetalle.query.delete()
        Producto.query.delete()
        Gasto.query.delete()
        AlertaStock.query.delete()
    else:
        return error('Tipo de limpieza no válido')
    db.session.commit()
    return jsonify({'success': True,
                    'mensaje': 'Todos los datos eliminados' if tipo == 'todo' else 'Historial de ventas eliminado'})

# ─── API: CONFIGURACIÓN Y CUENTA ─────────────────────────────────────────────

@app.route('/api/config', methods=['GET', 'POST'])
@login_required
def api_config():
    u = current_user
    if request.method == 'POST':
        d = request.get_json(silent=True) or {}
        if 'nombre' in d and texto(d['nombre']):
            u.nombre = texto(d['nombre'])[:100]
        if 'empresa' in d:
            u.empresa = texto(d['empresa'], 'Mi Empresa')[:100]
        if 'telefono' in d:
            u.telefono = texto(d['telefono'])[:50]
        for campo in ('capital_inicial', 'meta_mensual'):
            if campo in d:
                v = parse_num(d[campo], 0)
                if v is None or v < 0:
                    return error('Los valores no pueden ser negativos')
                setattr(u, campo, v)
        db.session.commit()
        return jsonify({'success': True})
    return jsonify({'nombre': u.nombre, 'email': u.email, 'empresa': u.empresa,
                    'capital_inicial': u.capital_inicial, 'meta_mensual': u.meta_mensual,
                    'rol': u.rol, 'telefono': u.telefono, 'debe_cambiar_clave': u.debe_cambiar_clave})


@app.route('/api/cuenta/clave', methods=['POST'])
@login_required
def api_cambiar_clave():
    d = request.get_json(silent=True) or {}
    actual, nueva = d.get('actual', ''), d.get('nueva', '')
    if not current_user.check_password(actual):
        return error('La contraseña actual no es correcta')
    if len(nueva) < 8:
        return error('La nueva contraseña debe tener al menos 8 caracteres')
    if nueva == actual:
        return error('La nueva contraseña debe ser diferente a la actual')
    current_user.set_password(nueva)
    cfg = current_user.get_config()
    cfg.pop('cambiar_clave', None)
    current_user.set_config(cfg)
    db.session.commit()
    return jsonify({'success': True})


@app.route('/api/usuarios', methods=['GET'])
@login_required
@rol_requerido('admin')
def api_usuarios():
    return jsonify([{'id': u.id, 'nombre': u.nombre, 'email': u.email, 'rol': u.rol}
                    for u in Usuario.query.order_by(Usuario.id).all()])


@app.route('/api/usuarios', methods=['POST'])
@login_required
@rol_requerido('admin')
def api_usuarios_crear():
    d = request.get_json(silent=True) or {}
    email = texto(d.get('email')).lower()
    nombre = texto(d.get('nombre'))
    clave = d.get('clave', '')
    rol = d.get('rol', 'vendedor')
    if not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', email):
        return error('Correo no válido')
    if not nombre:
        return error('Escribe el nombre')
    if len(clave) < 8:
        return error('La contraseña debe tener al menos 8 caracteres')
    if rol not in ROLES:
        return error('Rol no válido')
    if Usuario.query.filter(func.lower(Usuario.email) == email).first():
        return error('Ya existe un usuario con ese correo')
    u = Usuario(nombre=nombre[:100], email=email, rol=rol, empresa=current_user.empresa,
                meta_mensual=current_user.meta_mensual, capital_inicial=current_user.capital_inicial)
    u.set_password(clave)
    u.set_config({'cambiar_clave': True})
    db.session.add(u)
    db.session.commit()
    return jsonify({'success': True, 'id': u.id})


@app.route('/api/usuarios/<int:uid>', methods=['DELETE'])
@login_required
@rol_requerido('admin')
def api_usuarios_eliminar(uid):
    if uid == current_user.id:
        return error('No puedes eliminar tu propio usuario')
    u = db.session.get(Usuario, uid)
    if not u:
        return error('Usuario no encontrado', 404)
    db.session.delete(u)
    db.session.commit()
    return jsonify({'success': True})

# ─── INICIALIZACIÓN ──────────────────────────────────────────────────────────

def init_db():
    """Crea las tablas y, si no hay usuarios, el administrador inicial."""
    db.create_all()
    if Usuario.query.first():
        return
    email = os.environ.get('ADMIN_EMAIL', 'admin@inventa.com').strip().lower()
    clave = os.environ.get('ADMIN_PASSWORD', '')
    admin = Usuario(nombre=os.environ.get('ADMIN_NOMBRE', 'Admin'), email=email,
                    empresa=os.environ.get('EMPRESA', 'Mi Empresa'), rol='admin')
    if clave:
        admin.set_password(clave)
    else:
        # Sin ADMIN_PASSWORD se usa la clave de siempre, pero se obliga a cambiarla al entrar.
        admin.set_password('inventa2024')
        admin.set_config({'cambiar_clave': True})
    db.session.add(admin)
    db.session.commit()
    print(f'[OK] Usuario administrador creado: {email}')


with app.app_context():
    init_db()
    print(f"[INFO] Base de datos: {'SQLite local' if USA_SQLITE else 'Postgres'}"
          f" · modo {'producción' if ENV_PRODUCCION else 'local'}")
    if ENV_PRODUCCION and USA_SQLITE:
        print('[WARN] En producción con SQLite los datos se BORRAN en cada reinicio. Configura DATABASE_URL.')


if __name__ == '__main__':
    app.run(debug=os.environ.get('FLASK_DEBUG') == '1', host='127.0.0.1', port=5000)
